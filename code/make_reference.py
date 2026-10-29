"""Build the Word reference document for the report: pandoc's own reference, restyled to the
look of the course's IEEE template (ASSIGN_template.doc).

Why not the template itself: converted from .doc, its Normal style is centred, it lacks the
styles pandoc writes (Compact, First Paragraph, Source Code, Table), and its heading
numbering shares a list with pandoc's numbered lists, which printed runs of empty headings.
So the report types its section numbers (I., A.) and takes only the IEEE look from here:
A4 page and margins from the template, Times New Roman, 10 pt justified body, small-caps
centred section headings, italic subsection headings.

Run: uv run --with python-docx python code/make_reference.py
"""

import subprocess
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, Twips

HERE = Path(__file__).resolve().parent.parent
OUT = HERE / "build" / "ieee-reference.docx"


def main():
    OUT.parent.mkdir(exist_ok=True)
    subprocess.run(
        ["pandoc", "-o", str(OUT), "--print-default-data-file", "reference.docx"],
        check=True,
    )
    doc = Document(OUT)
    sec = doc.sections[0]  # the template's page: A4, margins in twips
    sec.page_width, sec.page_height = Twips(11906), Twips(16838)
    sec.left_margin = sec.right_margin = Twips(893)
    sec.top_margin, sec.bottom_margin = Twips(1080), Twips(1440)

    st = doc.styles
    for name in (
        "Normal",
        "Body Text",
        "First Paragraph",
        "Compact",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "Title",
        "Author",
        "Image Caption",
        "Table Caption",
        "Bibliography",
    ):
        try:
            style = st[name]
        except KeyError:
            continue
        style.font.name = "Times New Roman"
        style.font.color.rgb = None
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts"
        )
        if rfonts is not None:  # pandoc's theme fonts outrank the name unless cleared
            for attr in list(rfonts.attrib):
                if "Theme" in attr:
                    del rfonts.attrib[attr]
            for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
                rfonts.set(
                    f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{attr}",
                    "Times New Roman",
                )

    for name in ("Normal", "Body Text", "First Paragraph"):
        st[name].font.size = Pt(10)
    for name in ("Body Text", "First Paragraph"):
        pf = st[name].paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.first_line_indent = Pt(14)
        pf.space_before, pf.space_after = Pt(0), Pt(6)
    st["First Paragraph"].paragraph_format.first_line_indent = Pt(0)
    st["Compact"].font.size = Pt(10)
    # Compact sits on Body Text: without this, table cells and list items inherit its justification
    # and first-line indent, which wrapped narrow headers ("Scenari/o") and pushed cells right.
    cpf = st["Compact"].paragraph_format
    cpf.alignment, cpf.first_line_indent = WD_ALIGN_PARAGRAPH.LEFT, Pt(0)
    # keep-with-next on every Compact paragraph holds each (short) table on one page and a list
    # with the paragraph after it; the block after a table must not keep with its own next block,
    # or a table, a paragraph and a tall figure chain into one unbreakable run.
    cpf.keep_with_next = True
    # "Steps followed:" stays with its list; First Paragraph itself must not keep with the next
    # block, or a paragraph after a table would chain to a tall figure and leave half a page blank.
    lead = st.add_style("Steps Lead", WD_STYLE_TYPE.PARAGRAPH)
    lead.base_style = st["First Paragraph"]
    lead.paragraph_format.keep_with_next = True
    for n in ("Body Text", "First Paragraph"):
        st[n].paragraph_format.space_before = Pt(4)  # keeps text off the table above it

    h1 = st["Heading 1"]
    h1.font.size, h1.font.bold, h1.font.small_caps = Pt(11), False, True
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before, h1.paragraph_format.space_after = Pt(12), Pt(6)
    h2 = st["Heading 2"]
    h2.font.size, h2.font.bold, h2.font.italic = Pt(10), False, True
    h2.paragraph_format.space_before, h2.paragraph_format.space_after = Pt(8), Pt(4)

    title = st["Title"]
    title.font.size, title.font.bold = Pt(22), False
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author = st["Author"]
    author.font.size = Pt(11)
    author.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for name in ("Image Caption", "Table Caption"):
        cap = st[name]
        cap.font.size, cap.font.italic = Pt(8), False
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(10)
    st["Table Caption"].font.small_caps = True
    # a caption right after a list item otherwise touches it
    st["Table Caption"].paragraph_format.space_before = Pt(8)
    # a caption left alone at the foot of a page is read as the caption of nothing
    st["Table Caption"].paragraph_format.keep_with_next = True
    # the reference list (a Bibliography div in the report): left-aligned, since justifying a
    # line that holds one long URL stretches its spaces, with a hanging indent under the [n]
    bib = st["Bibliography"].paragraph_format
    bib.alignment = WD_ALIGN_PARAGRAPH.LEFT
    bib.left_indent, bib.first_line_indent = Pt(18), Pt(-18)
    bib.space_before, bib.space_after = Pt(0), Pt(4)
    st["Bibliography"].font.size = Pt(9)

    # inline code set at body size looks larger than Times New Roman beside it
    for name in ("Verbatim Char",):
        if name in [s.name for s in st]:
            st[name].font.size = Pt(9)
    code = st["Source Code"] if "Source Code" in [s.name for s in st] else None
    if code is not None:
        code.font.size = Pt(8)
    fig = st[
        "Captioned Figure"
    ]  # pandoc's figure paragraph is left-aligned, its caption centred
    fig.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fig.paragraph_format.keep_with_next = True
    fig.paragraph_format.space_before = Pt(8)

    # Table style: a visible grid, tighter cell margins so short headers stay on one line.
    tbl = st["Table"].element
    tblpr = tbl.find(qn("w:tblPr"))
    mar = tblpr.find(qn("w:tblCellMar"))
    for side in ("left", "right"):
        mar.find(qn("w:" + side)).set(qn("w:w"), "57")
    borders = parse_xml(
        f"<w:tblBorders {nsdecls('w')}>"
        + "".join(
            f'<w:{e} w:val="single" w:sz="4" w:space="0" w:color="808080"/>'
            for e in ("top", "left", "bottom", "right", "insideH", "insideV")
        )
        + "</w:tblBorders>"
    )
    tblpr.insert(0, borders)
    # the header row of a table is bold, so a reader tells it from the data at a glance
    tbl.append(
        parse_xml(
            f'<w:tblStylePr {nsdecls("w")} w:type="firstRow"><w:rPr><w:b/></w:rPr></w:tblStylePr>'
        )
    )

    # page numbers, centred in the footer: a reader of a 20-page PDF refers to a page
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
        if kind:
            run._r.append(
                parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="{kind}"/>')
            )
        else:
            instr = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"/>')
            instr.text = text
            run._r.append(instr)
    run.font.size = Pt(9)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
