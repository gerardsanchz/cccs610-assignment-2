"""Let the block after each table break onto the next page on its own.

Every table cell is a Compact paragraph, and Compact keeps with the next paragraph so a table
stays whole on one page. The last row would then chain the table to whatever follows it (a
paragraph, then a tall figure), which moves the whole run to the next page and leaves the foot
of the page empty. Switching keep-with-next off in the last row ends the chain at the table.

Usage: uv run --with python-docx python code/free_table_tails.py build/assignment-2.docx
"""

import sys

from docx import Document


def main(path: str) -> int:
    doc = Document(path)
    for table in doc.tables:
        for cell in table.rows[-1].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = False
    doc.save(path)
    print(f"{path}: freed the last row of {len(doc.tables)} tables")
    return 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    sys.exit(main(sys.argv[1]))
